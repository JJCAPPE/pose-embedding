#!/usr/bin/env python3
"""Build the Fall 2026 BU UROP application working draft."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ARCHIVE_DIR = Path(__file__).resolve().parent
OUTPUT_PATH = ARCHIVE_DIR / "BU_UROP_Fall_2026_Application_Draft.docx"

BLUE = "1E4D7B"
DARK_BLUE = "163653"
PALE_BLUE = "EAF2F8"
PALE_GOLD = "FFF5D9"
PALE_GREEN = "EAF5EF"
LIGHT_GRAY = "F3F5F7"
MID_GRAY = "66727D"
DARK = "1F2933"
WHITE = "FFFFFF"


TITLE = "Comparative Metric Learning for Robust Pose-Sequence Retrieval"

DESCRIPTION = [
    (
        "Human-motion retrieval systems map a sequence of body-joint coordinates "
        "to an embedding so that motions with the same semantics appear near one "
        "another. This supports nearest-neighbor search for action recognition, "
        "technique comparison, and retrieval of reference movements. Unlike "
        "ordinary image embeddings, pose sequences are usually estimated from "
        "video and therefore contain structured errors: joint jitter, occlusion, "
        "missing limbs or frames, timing noise, and estimator-dependent bias."
    ),
    (
        "I will test whether contextual metric learning, introduced by Liao, "
        "Tsiligkaridis, and Kulis for image retrieval [1], transfers to "
        "pose-sequence retrieval. Contextual loss does not optimize only the "
        "similarity of individual pairs; it also encourages examples that should "
        "match to share local neighborhoods. My core hypothesis is that this "
        "neighborhood constraint will preserve ranking quality more effectively "
        "than direct pairwise or tuple losses as pose observations become "
        "corrupted. This is an untested input-noise question: the 2023 paper "
        "established robustness to label noise in images, not to corrupted "
        "skeleton coordinates."
    ),
    (
        "During Fall 2026 I will (1) build a reproducible retrieval pipeline using "
        "NTU RGB+D 120 skeleton sequences [3] and a fixed pretrained MotionBERT "
        "encoder [2]; (2) validate a clean baseline and compare supervised "
        "contrastive, triplet, Multi-Similarity, contextual-only, and the "
        "published contextual-plus-contrastive objective; (3) evaluate the same "
        "models on deterministic pose corruptions while keeping a clean gallery; "
        "and (4) begin a paper-fidelity transfer of representative pairwise, "
        "proxy, and rank-optimization objectives from the 2023 study without "
        "changing the encoder. Planned deliverables are tested code, fixed data "
        "manifests, clean/noisy Recall@K and mean-average-precision curves, and an "
        "analysis of robustness, runtime, and failure modes. If the core matrix is "
        "stable, I will pilot one recent compatible method rather than broaden the "
        "encoder or dataset. My individual role is to implement the pipeline, run "
        "the experiments, analyze the results, and prepare the report under "
        "Prof. Kulis's supervision."
    ),
]

SIGNIFICANCE = [
    (
        "Pose embeddings can make large motion collections searchable: a new "
        "sequence can retrieve actions or reference movements with similar "
        "structure instead of requiring an exact label or a hand-designed "
        "distance. This capability is relevant to video archives, sports "
        "technique analysis, and rehabilitation research, but only if retrieval "
        "remains reliable when the input skeleton is incomplete or noisy."
    ),
    (
        "The scientific gap is not simply the absence of another motion model. "
        "Contextual similarity and many competing metric-learning objectives were "
        "developed and ranked on images. Pose sequences have different geometry: "
        "joints are anatomically coupled, time matters, similarity may be graded, "
        "and errors affect connected body parts across frames. Consequently, an "
        "objective that performs well for clean images may not transfer—or may "
        "fail differently—on corrupted motion data. The contextual paper studied "
        "label noise [1], while recent skeleton work has shown that small pose "
        "perturbations and changes of pose estimator can substantially degrade "
        "learned representations [8]."
    ),
    (
        "This project will isolate the optimization objective by holding the "
        "encoder, embedding dimension, data splits, training budget, and evaluator "
        "fixed. That control is important because published metric-learning "
        "comparisons often mix losses with different backbones, miners, or "
        "training recipes [4]. The result will show whether neighborhood context, "
        "pair mining, class proxies, or direct rank optimization best preserves "
        "pose retrieval under controlled corruption."
    ),
    (
        "Either outcome is informative. A positive result would identify a "
        "robustness advantage for contextual training; a null result would show "
        "that a simpler objective is sufficient under a fair budget; and a "
        "reversal of the image-domain ranking would establish that objective gains "
        "are domain-dependent. The code, corruption protocol, and portability "
        "analysis will provide a reproducible basis for later fine-grained motion "
        "or real-video studies without making clinical or performance claims in "
        "this semester."
    ),
]

METHODOLOGY = [
    (
        "Data and split.",
        "I will use the existing NTU RGB+D 120 skeleton dataset [3]. The standard "
        "cross-subject protocol will separate training and test identities; I "
        "will derive a subject-disjoint validation partition from the training "
        "side. Test samples will be assigned once to disjoint query and gallery "
        "manifests, with action class defining relevance. Self-matches and duplicate "
        "sequence leakage will be excluded."
    ),
    (
        "Representation.",
        "After root-centering, scale normalization, and deterministic temporal "
        "handling, a pretrained MotionBERT/DSTFormer encoder [2] will produce "
        "sequence features. I will first freeze the encoder, use one fixed pooling "
        "rule, and train the same small projection head for every objective. A "
        "frozen-feature baseline will verify the evaluator before training. Full "
        "encoder fine-tuning will be attempted only if compute permits and then "
        "applied identically across compared objectives."
    ),
    (
        "Controlled objectives.",
        "The core comparison will include supervised contrastive loss [5], triplet "
        "loss with its miner, Multi-Similarity with and without mining [6], "
        "contextual-only loss, and the published contextual objective "
        "(contextual + contrastive + similarity regularization) [1]. Once these "
        "pass deterministic smoke tests, I will add representative objectives from "
        "the 2023 paper-fidelity suite: contrastive, Proxy Anchor, Proxy NCA, and "
        "ROADMAP. The backbone, initialization, projection dimension, batches, "
        "optimizer family, horizon, and seeds will remain fixed. Each loss will "
        "receive the same number of validation trials, rather than inappropriate "
        "identical margin values."
    ),
    (
        "Corruption benchmark.",
        "Clean and corrupted versions of each query will be evaluated against the "
        "same clean gallery. Seeded operators will apply joint-coordinate jitter "
        "scaled to skeleton size, missing-joint or limb masks, frame dropout, and "
        "temporal jitter at several severity levels. Unit tests will verify that "
        "labels, gallery items, and non-target joints remain unchanged. If the core "
        "comparison is complete, I will pilot one compatible post-2023 method or a "
        "pose-specific robustness control such as MaskCLR [8]."
    ),
    (
        "Evaluation and analysis.",
        "Cosine similarity will rank gallery embeddings. I will report Recall@1, "
        "Recall@5, Recall@10, mean average precision, mAP@R, absolute degradation "
        "from clean performance, normalized retention, and area under each "
        "corruption-severity curve. Paired corruption seeds and repeated training "
        "seeds will support mean/standard-deviation reporting. I will also record "
        "training time, peak memory, convergence failures, and method-specific "
        "parameters. Fixed manifests, configuration files, unit tests, and a "
        "results table generated from saved outputs will make the comparison "
        "auditable."
    ),
]

TIMELINE = [
    (
        "Week 1 — Protocol and access: finalize the research question with "
        "Prof. Kulis; confirm NTU RGB+D 120 access, compute resources, and the "
        "subject-disjoint split; freeze success criteria and configuration schema."
    ),
    (
        "Week 2 — Data pipeline: implement preprocessing, sequence masks or "
        "resampling, deterministic train/validation/test manifests, and a small "
        "debug subset; add leakage and shape tests."
    ),
    (
        "Week 3 — Retrieval baseline: extract frozen MotionBERT features; implement "
        "cosine query/gallery evaluation and Recall@K, mAP, and mAP@R; verify "
        "self-match exclusion and reproduce results from saved embeddings."
    ),
    (
        "Week 4 — Corruptions: implement seeded joint jitter, missing joints/limbs, "
        "frame dropout, and temporal jitter; validate severity scaling and generate "
        "clean-versus-corrupted baseline curves."
    ),
    (
        "Week 5 — Core objectives: train the shared projection head with supervised "
        "contrastive, triplet, and Multi-Similarity objectives under matched "
        "manifests and validation budgets."
    ),
    (
        "Week 6 — Contextual objective: validate contextual-only and full contextual "
        "training, including balanced-batch requirements, gradient checks, and "
        "component ablations."
    ),
    (
        "Week 7 — Core comparison: run paired clean/corrupted evaluations across "
        "multiple seeds; profile time and memory; diagnose unstable or unfair "
        "settings before expansion."
    ),
    (
        "Week 8 — Transfer extension: add one proxy objective and one rank-oriented "
        "objective from the 2023 comparison. If the core matrix is not stable, use "
        "this week for replication and ablation instead of adding methods."
    ),
    (
        "Week 9 — Analysis: aggregate mean and variability, robustness retention, "
        "severity-curve area, and compute cost; compare the pose-domain method "
        "ordering with the image-domain evidence and document limitations."
    ),
    (
        "Week 10 — Deliverables: complete the reproducible experiment package, "
        "technical report, final figures/tables, and a short presentation or poster; "
        "review conclusions and next-stage objective priorities with Prof. Kulis."
    ),
]

BACKGROUND = [
    (
        "I am a Boston University Computer Engineering student with completed "
        "coursework in Computational Linear Algebra (ENG EK 103), Probability, "
        "Statistics, and Data Science (ENG EK 381), Introduction to Software "
        "Engineering (ENG EC 327), multivariable calculus, and differential "
        "equations. This preparation is directly relevant to vector-space "
        "representations, stochastic evaluation, optimization, and reproducible "
        "software."
    ),
    (
        "From January through May 2026, I worked as an undergraduate machine-"
        "learning researcher in the BU College of Engineering on a video-first "
        "rowing-biomechanics pipeline. I built Python components for 2D pose "
        "extraction with MMPose/Sports2D, MotionBERT-based 3D motion "
        "representations, temporal alignment of kinematics with force telemetry, "
        "quality-control checks, and sequence-to-sequence regression. That work "
        "gave me direct experience with the exact failure modes motivating this "
        "project: occlusion, joint jitter, camera-dependent error, timing drift, "
        "and imperfect 2D-to-3D lifting."
    ),
    (
        "My technical experience includes Python, PyTorch, NumPy, SciPy, "
        "scikit-learn, pandas, Git, automated tests, and experiment/report "
        "pipelines. In preparation for this project, I have reviewed the "
        "contextual-similarity paper and its official implementation, surveyed "
        "pose and metric-learning baselines, and assembled an auditable PyTorch "
        "reference implementation with deterministic gradient and shape checks."
    ),
    (
        "I developed this proposal after discussing the research direction with "
        "Prof. Kulis. His feedback led to the controlled extension: keep the pose "
        "encoder fixed, first establish the contextual-loss result, and then map "
        "the 2023 comparison objectives and selected newer methods into the pose "
        "domain. The project advances my goal of conducting rigorous machine-"
        "learning and computer-vision research while building on a demonstrated "
        "ability to implement and evaluate pose-processing systems."
    ),
]

BIBLIOGRAPHY = [
    (
        "Liao, C., Tsiligkaridis, T., & Kulis, B. (2023). Supervised Metric "
        "Learning to Rank for Retrieval via Contextual Similarity Optimization. "
        "ICML, PMLR 202:20906–20938. "
        "https://proceedings.mlr.press/v202/liao23b.html"
    ),
    (
        "Zhu, W., Ma, X., Liu, Z., Liu, L., Wu, W., & Wang, Y. (2023). MotionBERT: "
        "A Unified Perspective on Learning Human Motion Representations. ICCV. "
        "https://arxiv.org/abs/2210.06551"
    ),
    (
        "Liu, J., Shahroudy, A., Perez, M., Wang, G., Duan, L.-Y., & Kot, A. C. "
        "(2020). NTU RGB+D 120: A Large-Scale Benchmark for 3D Human Activity "
        "Understanding. IEEE TPAMI, 42(10), 2684–2701. "
        "https://doi.org/10.1109/TPAMI.2019.2916873"
    ),
    (
        "Musgrave, K., Belongie, S., & Lim, S.-N. (2020). A Metric Learning Reality "
        "Check. ECCV. https://arxiv.org/abs/2003.08505"
    ),
    (
        "Khosla, P., Teterwak, P., Wang, C., et al. (2020). Supervised Contrastive "
        "Learning. NeurIPS 33. https://arxiv.org/abs/2004.11362"
    ),
    (
        "Wang, X., Han, X., Huang, W., Dong, D., & Scott, M. R. (2019). "
        "Multi-Similarity Loss with General Pair Weighting for Deep Metric "
        "Learning. CVPR, 5022–5030."
    ),
    (
        "Kim, S., Kim, D., Cho, M., & Kwak, S. (2020). Proxy Anchor Loss for Deep "
        "Metric Learning. CVPR, 3238–3247."
    ),
    (
        "Abdelfattah, M., Hassan, M., & Alahi, A. (2024). MaskCLR: Attention-Guided "
        "Contrastive Learning for Robust Action Representation Learning. CVPR, "
        "18678–18687. https://doi.org/10.1109/CVPR52733.2024.01767"
    ),
    (
        "Kim, S., Seo, M., Laptev, I., Cho, M., & Kwak, S. (2019). Deep Metric "
        "Learning Beyond Binary Supervision. CVPR, 2288–2297. "
        "https://openaccess.thecvf.com/content_CVPR_2019/html/"
        "Kim_Deep_Metric_Learning_Beyond_Binary_Supervision_CVPR_2019_paper.html"
    ),
]


def plain_text(paragraphs: list[str]) -> str:
    return "\n\n".join(paragraphs)


def labeled_plain_text(items: list[tuple[str, str]]) -> str:
    return "\n\n".join(f"{label} {body}" for label, body in items)


def list_plain_text(items: list[str]) -> str:
    return "\n".join(f"• {item}" for item in items)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top: int = 90, start: int = 110, bottom: int = 90, end: int = 110) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_table_cell_width(cell, width_inches: float) -> None:
    width = Inches(width_inches)
    cell.width = width
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width)))
    tc_w.set(qn("w:type"), "dxa")


def shade_paragraph(paragraph, fill: str, border_color: str | None = None) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    shd = p_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        p_pr.append(shd)
    shd.set(qn("w:fill"), fill)
    if border_color:
        p_bdr = p_pr.find(qn("w:pBdr"))
        if p_bdr is None:
            p_bdr = OxmlElement("w:pBdr")
            p_pr.append(p_bdr)
        left = OxmlElement("w:left")
        left.set(qn("w:val"), "single")
        left.set(qn("w:sz"), "18")
        left.set(qn("w:space"), "8")
        left.set(qn("w:color"), border_color)
        p_bdr.append(left)


def set_paragraph_padding(paragraph, before: int = 80, after: int = 80) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    spacing = p_pr.find(qn("w:spacing"))
    if spacing is None:
        spacing = OxmlElement("w:spacing")
        p_pr.append(spacing)
    spacing.set(qn("w:before"), str(before))
    spacing.set(qn("w:after"), str(after))


def add_field(run, instruction: str) -> None:
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])


def add_hyperlink(paragraph, text: str, url: str, color: str = BLUE) -> None:
    part = paragraph.part
    relationship_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    run_pr = OxmlElement("w:rPr")
    run_color = OxmlElement("w:color")
    run_color.set(qn("w:val"), color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    run_pr.extend([run_color, underline])
    run.append(run_pr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_text_with_urls(paragraph, text: str, *, bold_prefix: str | None = None) -> None:
    if bold_prefix and text.startswith(bold_prefix):
        run = paragraph.add_run(bold_prefix)
        run.bold = True
        text = text[len(bold_prefix):]
    url_re = re.compile(r"(https?://\S+)")
    position = 0
    for match in url_re.finditer(text):
        if match.start() > position:
            paragraph.add_run(text[position:match.start()])
        url = match.group(1).rstrip(".,);")
        trailing = match.group(1)[len(url):]
        add_hyperlink(paragraph, url, url)
        if trailing:
            paragraph.add_run(trailing)
        position = match.end()
    if position < len(text):
        paragraph.add_run(text[position:])


def style_document(document: Document) -> None:
    styles = document.styles

    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(DARK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    normal.paragraph_format.line_spacing = 1.25

    for style_name in ("List Bullet", "List Number"):
        style = styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(10.5)
        style.font.color.rgb = RGBColor.from_string(DARK)
        style.paragraph_format.space_after = Pt(3)
        style.paragraph_format.line_spacing = 1.16

    heading_1 = styles["Heading 1"]
    heading_1.font.name = "Calibri"
    heading_1.font.size = Pt(16)
    heading_1.font.bold = True
    heading_1.font.color.rgb = RGBColor.from_string(BLUE)
    heading_1.paragraph_format.space_before = Pt(16)
    heading_1.paragraph_format.space_after = Pt(8)
    heading_1.paragraph_format.keep_with_next = True

    heading_2 = styles["Heading 2"]
    heading_2.font.name = "Calibri"
    heading_2.font.size = Pt(13)
    heading_2.font.bold = True
    heading_2.font.color.rgb = RGBColor.from_string(BLUE)
    heading_2.paragraph_format.space_before = Pt(12)
    heading_2.paragraph_format.space_after = Pt(6)
    heading_2.paragraph_format.keep_with_next = True

    heading_3 = styles["Heading 3"]
    heading_3.font.name = "Calibri"
    heading_3.font.size = Pt(12)
    heading_3.font.bold = True
    heading_3.font.color.rgb = RGBColor.from_string(DARK_BLUE)
    heading_3.paragraph_format.space_before = Pt(8)
    heading_3.paragraph_format.space_after = Pt(4)
    heading_3.paragraph_format.keep_with_next = True


def configure_page(section) -> None:
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.78)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)
    section.header_distance = Inches(0.32)
    section.footer_distance = Inches(0.34)


def add_header_footer(section) -> None:
    header = section.header
    header.is_linked_to_previous = False
    paragraph = header.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_after = Pt(0)
    left = paragraph.add_run("BU UROP  ·  FALL 2026")
    left.font.name = "Calibri"
    left.font.size = Pt(8.5)
    left.font.bold = True
    left.font.color.rgb = RGBColor.from_string(BLUE)
    paragraph.add_run("\t")
    right = paragraph.add_run("WORKING DRAFT  ·  31 JULY 2026")
    right.font.name = "Calibri"
    right.font.size = Pt(8.5)
    right.font.color.rgb = RGBColor.from_string(MID_GRAY)
    tabs = paragraph.paragraph_format.tab_stops
    tabs.add_tab_stop(Inches(6.6))

    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), "B8C9D8")
    p_bdr.append(bottom)
    p_pr.append(p_bdr)

    footer = section.footer
    footer.is_linked_to_previous = False
    footer_p = footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_p.paragraph_format.space_before = Pt(0)
    run = footer_p.add_run("Giacomo Cappelletto  ·  ")
    run.font.name = "Calibri"
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string(MID_GRAY)
    page_run = footer_p.add_run()
    page_run.font.name = "Calibri"
    page_run.font.size = Pt(8)
    page_run.font.color.rgb = RGBColor.from_string(MID_GRAY)
    add_field(page_run, "PAGE")


def add_label_value(document: Document, label: str, value: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(2)
    run = paragraph.add_run(f"{label}  ")
    run.bold = True
    run.font.color.rgb = RGBColor.from_string(DARK_BLUE)
    paragraph.add_run(value)


def add_callout(document: Document, title: str, body: str, fill: str, border: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.08)
    paragraph.paragraph_format.right_indent = Inches(0.05)
    set_paragraph_padding(paragraph, before=100, after=100)
    shade_paragraph(paragraph, fill, border)
    title_run = paragraph.add_run(f"{title}\n")
    title_run.bold = True
    title_run.font.color.rgb = RGBColor.from_string(DARK_BLUE)
    paragraph.add_run(body)


def add_source_line(document: Document, label: str, url: str, suffix: str = "") -> None:
    paragraph = document.add_paragraph(style="List Bullet")
    run = paragraph.add_run(f"{label}: ")
    run.bold = True
    add_hyperlink(paragraph, url, url)
    if suffix:
        paragraph.add_run(f" — {suffix}")


def add_char_count(document: Document, text: str, limit: int) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(1)
    paragraph.paragraph_format.space_after = Pt(5)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run(
        f"{len(text):,} / {limit:,} characters, including spaces"
    )
    run.font.name = "Calibri"
    run.font.size = Pt(8.5)
    run.font.bold = True
    run.font.color.rgb = RGBColor.from_string(MID_GRAY)


def add_application_section(
    document: Document,
    heading: str,
    paragraphs: list[str],
    limit: int = 3000,
) -> None:
    document.add_heading(heading, level=2)
    text = plain_text(paragraphs)
    add_char_count(document, text, limit)
    for body in paragraphs:
        paragraph = document.add_paragraph()
        add_text_with_urls(paragraph, body)


def add_labeled_application_section(
    document: Document,
    heading: str,
    items: list[tuple[str, str]],
    limit: int = 3000,
) -> None:
    document.add_heading(heading, level=2)
    text = labeled_plain_text(items)
    add_char_count(document, text, limit)
    for label, body in items:
        paragraph = document.add_paragraph()
        run = paragraph.add_run(f"{label} ")
        run.bold = True
        run.font.color.rgb = RGBColor.from_string(DARK_BLUE)
        paragraph.add_run(body)


def add_timeline_section(document: Document) -> None:
    document.add_heading("Timeline", level=2)
    text = list_plain_text(TIMELINE)
    add_char_count(document, text, 3000)
    for item in TIMELINE:
        paragraph = document.add_paragraph(style="List Bullet")
        week_label, separator, body = item.partition(" — ")
        if separator:
            run = paragraph.add_run(f"{week_label} — ")
            run.bold = True
            run.font.color.rgb = RGBColor.from_string(DARK_BLUE)
            paragraph.add_run(body)
        else:
            paragraph.add_run(item)


def add_bibliography(document: Document) -> None:
    document.add_heading("Bibliography", level=2)
    text = "\n".join(f"{index}. {item}" for index, item in enumerate(BIBLIOGRAPHY, 1))
    add_char_count(document, text, 3000)
    for item in BIBLIOGRAPHY:
        paragraph = document.add_paragraph(style="List Number")
        paragraph.paragraph_format.left_indent = Inches(0.27)
        paragraph.paragraph_format.first_line_indent = Inches(-0.18)
        add_text_with_urls(paragraph, item)


def add_funding_table(document: Document) -> None:
    table = document.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.style = "Table Grid"
    widths = [1.15, 1.0, 1.3, 3.0]
    headers = ["Option", "Hours/week", "Award", "Funding source"]
    for index, (cell, width, header) in enumerate(zip(table.rows[0].cells, widths, headers)):
        set_table_cell_width(cell, width)
        set_cell_shading(cell, BLUE)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if index in (1, 2) else WD_ALIGN_PARAGRAPH.LEFT
        run = paragraph.add_run(header)
        run.bold = True
        run.font.color.rgb = RGBColor.from_string(WHITE)
        run.font.size = Pt(9.5)
    set_repeat_table_header(table.rows[0])

    rows = [
        ("SRA", "5", "$750", "UROP provides the full amount"),
        ("FMG", "5", "$750", "$375 UROP + $375 mentor"),
        ("FMG", "8", "$1,200", "$600 UROP + $600 mentor"),
        ("FMG", "10", "$1,500", "$750 UROP + $750 mentor"),
    ]
    for row_index, row_data in enumerate(rows):
        cells = table.add_row().cells
        for index, (cell, width, value) in enumerate(zip(cells, widths, row_data)):
            set_table_cell_width(cell, width)
            set_cell_margins(cell)
            if row_index % 2 == 1:
                set_cell_shading(cell, LIGHT_GRAY)
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if index in (1, 2) else WD_ALIGN_PARAGRAPH.LEFT
            run = paragraph.add_run(value)
            run.font.size = Pt(9.5)
            if index == 0:
                run.bold = True


def build_document() -> Document:
    document = Document()
    style_document(document)
    configure_page(document.sections[0])
    add_header_footer(document.sections[0])

    core = document.core_properties
    core.title = "BU UROP Fall 2026 Application Draft"
    core.subject = "Comparative metric learning for robust pose-sequence retrieval"
    core.author = "Giacomo Cappelletto"
    core.keywords = "BU UROP, metric learning, pose retrieval, MotionBERT, contextual similarity"

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(12)
    title.paragraph_format.space_after = Pt(3)
    run = title.add_run("BU UROP · FALL 2026")
    run.font.name = "Calibri"
    run.font.size = Pt(10)
    run.font.bold = True
    run.font.color.rgb = RGBColor.from_string(BLUE)

    project_title = document.add_paragraph()
    project_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    project_title.paragraph_format.space_after = Pt(3)
    run = project_title.add_run(TITLE)
    run.font.name = "Calibri"
    run.font.size = Pt(22)
    run.font.bold = True
    run.font.color.rgb = RGBColor.from_string(DARK_BLUE)

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(12)
    run = subtitle.add_run("Funding application working draft · verified 31 July 2026")
    run.font.name = "Calibri"
    run.font.size = Pt(11)
    run.font.italic = True
    run.font.color.rgb = RGBColor.from_string(MID_GRAY)

    add_label_value(document, "Applicant", "Giacomo Cappelletto · B.S. Computer Engineering, Boston University")
    add_label_value(
        document,
        "Faculty mentor",
        "Prof. Brian Kulis · Electrical & Computer Engineering; affiliated with CS, SE, and CDS",
    )
    add_label_value(document, "Funding period", "Fall 2026 · ten paid weeks")

    add_callout(
        document,
        "Three confirmations before submission",
        (
            "1) Choose SRA or FMG and confirm any faculty match with Prof. Kulis. "
            "2) Confirm whether you have previously received UROP funding. "
            "3) Obtain a BU IRB determination for the secondary analysis of an "
            "existing human-motion dataset before answering the human-subjects item."
        ),
        PALE_GOLD,
        "D39B27",
    )

    document.add_heading("Verified application facts", level=1)
    facts = [
        (
            "Student deadline",
            "Friday, September 11, 2026, at 12:00 p.m.; the mentor recommendation is due by midnight.",
        ),
        (
            "Components",
            "A student application and a confidential faculty-mentor recommendation are both required.",
        ),
        (
            "Eligibility",
            "Full-time BU undergraduate, BU faculty supervision, no simultaneous academic credit and UROP stipend for the same research, and no more than two prior funded terms under the current eligibility page.",
        ),
        (
            "Drafting format",
            "The currently linked sample lists 500 characters for the title and 3,000 characters for each main narrative field. Recheck the live Fall 2026 form when it opens in early August.",
        ),
        (
            "Mentor action",
            "BU pages conflict about automatic mentor notification. Send Prof. Kulis the faculty recommendation link directly and do not rely on an automated email.",
        ),
    ]
    for label, body in facts:
        paragraph = document.add_paragraph(style="List Bullet")
        run = paragraph.add_run(f"{label}: ")
        run.bold = True
        paragraph.add_run(body)

    source_paragraph = document.add_paragraph()
    source_paragraph.paragraph_format.space_before = Pt(4)
    source_run = source_paragraph.add_run("Official BU sources: ")
    source_run.bold = True
    links = [
        ("application", "https://www.bu.edu/urop/get-involved/application-2/"),
        ("deadlines", "https://www.bu.edu/urop/get-involved/deadlines/"),
        ("funding and eligibility", "https://www.bu.edu/urop/get-involved/funding-eligibility/"),
        ("application tips", "https://www.bu.edu/urop/get-involved/application-tips/"),
    ]
    for index, (label, url) in enumerate(links):
        if index:
            source_paragraph.add_run(" · ")
        add_hyperlink(source_paragraph, label, url)

    document.add_heading("Funding decision", level=1)
    add_funding_table(document)
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(5)
    paragraph.add_run("Scope recommendation. ").bold = True
    paragraph.add_run(
        "An 8-hour FMG best matches the comparative scope if Prof. Kulis confirms "
        "the $600 faculty match. If the application uses the 5-hour SRA, preserve "
        "the contextual/SupCon/triplet/Multi-Similarity core and treat the proxy, "
        "rank-loss, and recent-method experiments as decision-gated extensions."
    )

    add_callout(
        document,
        "Authorship and review",
        (
            "The UROP form requires you to certify that the application is your own "
            "work and that sources are acknowledged. Review every sentence, revise "
            "it into your voice, verify the live form, and have Prof. Kulis review "
            "the complete application at least once before submission."
        ),
        PALE_BLUE,
        BLUE,
    )

    document.add_heading("Copy-Paste Application Responses", level=1)
    intro = document.add_paragraph()
    intro.add_run("Use note. ").bold = True
    intro.add_run(
        "Counts below include spaces and the response text only. They leave "
        "substantial margin under the currently published limits; confirm each "
        "count in the live Fall 2026 form."
    )

    document.add_heading("Project Title", level=2)
    add_char_count(document, TITLE, 500)
    title_box = document.add_paragraph()
    title_box.paragraph_format.left_indent = Inches(0.12)
    title_box.paragraph_format.right_indent = Inches(0.12)
    title_box.paragraph_format.space_after = Pt(8)
    set_paragraph_padding(title_box, before=100, after=100)
    shade_paragraph(title_box, PALE_BLUE, BLUE)
    run = title_box.add_run(TITLE)
    run.bold = True
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor.from_string(DARK_BLUE)

    add_application_section(document, "Project Description and Goals", DESCRIPTION)
    add_application_section(document, "Project Significance / Importance", SIGNIFICANCE)
    add_labeled_application_section(document, "Methodology / Process", METHODOLOGY)
    add_timeline_section(document)
    add_application_section(document, "Background Experience", BACKGROUND)
    add_bibliography(document)

    document.add_heading("Funds Requested", level=1)
    paragraph = document.add_paragraph()
    paragraph.add_run("Select exactly one option in the live form after mentor confirmation.").bold = True

    document.add_heading("Option A — FMG, 8 hours/week", level=2)
    fmg_text = (
        "I request a $1,200 Faculty Matching Grant for approximately eight hours "
        "per week during the ten-week Fall 2026 award period. UROP would provide "
        "$600 and the faculty mentor would provide the required $600 match. I am "
        "not requesting separate supplies or travel funding; the project will use "
        "an existing research dataset and available computing resources."
    )
    add_char_count(document, fmg_text, 3000)
    document.add_paragraph(fmg_text)

    document.add_heading("Option B — SRA, 5 hours/week", level=2)
    sra_text = (
        "I request a $750 Student Research Award for approximately five hours per "
        "week during the ten-week Fall 2026 award period. I am not requesting "
        "separate supplies or travel funding; the project will use an existing "
        "research dataset and available computing resources."
    )
    add_char_count(document, sra_text, 3000)
    document.add_paragraph(sra_text)

    document.add_heading("Safety and Research Compliance", level=1)
    add_callout(
        document,
        "IRB determination required before selecting the form response",
        (
            "The planned work is a secondary computational analysis of an existing "
            "human-motion dataset. It involves no recruitment, interaction, "
            "intervention, or new collection by the student. That fact does not "
            "authorize an independent “no human subjects” determination. Contact "
            "the BU IRB with Prof. Kulis and retain the office's written decision "
            "or exemption documentation before submission or before beginning any "
            "work for which review is required."
        ),
        PALE_GOLD,
        "D39B27",
    )

    compliance_items = [
        ("Animal subjects", "No animals are planned."),
        (
            "Human subjects",
            "Answer only after written guidance from the BU IRB; upload the required approval or exemption documentation if applicable.",
        ),
        (
            "Safety training",
            "Computational work only; no chemical, biological, radiological, clinical, or fabrication procedures are planned. Complete any standard computing or data-governance training required by the mentor or dataset.",
        ),
        (
            "Data handling",
            "Use only the approved research dataset and BU-approved storage/compute; do not attempt re-identification or combine subject metadata beyond the approved protocol.",
        ),
    ]
    for label, body in compliance_items:
        paragraph = document.add_paragraph(style="List Bullet")
        run = paragraph.add_run(f"{label}: ")
        run.bold = True
        paragraph.add_run(body)

    document.add_heading("Administrative Preflight", level=1)
    checklist = [
        "Fill in BU ID, BU email, school/class year, and any other live-form applicant fields.",
        "Optional GPA: 3.97 on the unofficial transcript printed June 11, 2026; recheck before entering it.",
        "Confirm prior UROP funding status. If yes, draft both 1,500-character continuation fields and verify eligibility under the current two-award cap.",
        "Choose one stipend option; confirm the FMG match in writing if applicable.",
        "Confirm that no academic credit or duplicate pay will cover the same research hours.",
        "Check total on-campus employment against BU's 20-hour weekly academic-year limit.",
        "Obtain the IRB determination and any required document before answering the human-subjects questions.",
        "Paste each response into the live form and recheck its displayed character count.",
        "Send the complete draft to Prof. Kulis for review; revise after his feedback.",
        "Send Prof. Kulis the Kerberos-protected recommendation-form link directly.",
        "Submit the student application before noon on September 11, 2026; verify mentor submission before midnight.",
    ]
    for item in checklist:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.left_indent = Inches(0.24)
        paragraph.paragraph_format.first_line_indent = Inches(-0.18)
        paragraph.paragraph_format.space_after = Pt(3)
        run = paragraph.add_run("☐ ")
        run.font.color.rgb = RGBColor.from_string(BLUE)
        paragraph.add_run(item)

    document.add_heading("Known BU Source Conflicts", level=1)
    conflict_items = [
        (
            "Mentor notification",
            "The Application page says the system notifies the mentor; Application Tips says it does not. Send the link directly.",
        ),
        (
            "Appendix length",
            "Application Tips says two pages; the linked March 2025 sample says one. No appendix is recommended for this draft.",
        ),
        (
            "Prior funding cap",
            "The current eligibility page says two funded terms; the older sample says three. Use the current two-award rule unless UROP confirms otherwise.",
        ),
    ]
    for label, body in conflict_items:
        paragraph = document.add_paragraph(style="List Bullet")
        run = paragraph.add_run(f"{label}: ")
        run.bold = True
        paragraph.add_run(body)

    closing = document.add_paragraph()
    closing.paragraph_format.space_before = Pt(8)
    closing.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = closing.add_run("End of working draft")
    run.font.size = Pt(9)
    run.font.italic = True
    run.font.color.rgb = RGBColor.from_string(MID_GRAY)

    return document


def main() -> None:
    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    document = build_document()
    document.save(OUTPUT_PATH)

    counts = {
        "title": len(TITLE),
        "description": len(plain_text(DESCRIPTION)),
        "significance": len(plain_text(SIGNIFICANCE)),
        "methodology": len(labeled_plain_text(METHODOLOGY)),
        "timeline": len(list_plain_text(TIMELINE)),
        "background": len(plain_text(BACKGROUND)),
        "bibliography": len(
            "\n".join(f"{index}. {item}" for index, item in enumerate(BIBLIOGRAPHY, 1))
        ),
    }
    print(f"Wrote {OUTPUT_PATH}")
    for field, count in counts.items():
        print(f"{field}: {count}")


if __name__ == "__main__":
    main()
